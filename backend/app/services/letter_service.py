"""Letter generation service."""

import os
from pathlib import Path
from docx import Document
from sqlmodel import Session

from app.exceptions import NotFoundException
from app.models import Letter
from app.services.approval_service import require_approved


def generate_docx(
    session: Session,
    letter_id: int,
    tenant_id: int,
) -> Letter:
    """Generate DOCX for letter, requiring offer_letter approval first.
    
    Saves file to backend/storage/<tenant_id>/<letter_id>.docx
    and updates the letter record with file_path.
    """
    letter = session.get(Letter, letter_id)
    if not letter or letter.tenant_id != tenant_id:
        raise NotFoundException(f"Letter {letter_id} not found")

    # Verify offer_letter approval is approved
    if letter.approval_id is None:
        raise NotFoundException("Letter has no associated offer approval")

    require_approved(
        session=session,
        approval_id=letter.approval_id,
        tenant_id=letter.tenant_id,
        expected_action_type="offer_letter",
    )

    # Build storage path - prevent path traversal by using only IDs
    storage_dir = Path("backend") / "storage" / str(letter.tenant_id)
    storage_dir.mkdir(parents=True, exist_ok=True)

    # Use only letter_id in filename to prevent path traversal
    file_name = f"{letter_id}.docx"
    file_path = storage_dir / file_name

    # Create DOCX document
    doc = Document()
    if letter.content:
        # Split content into subject and body
        lines = letter.content.split("\n")
        if lines and lines[0].startswith("Subject: "):
            subject = lines[0][9:]  # Remove "Subject: "
            body_lines = lines[2:] if len(lines) > 2 else []
            body = "\n".join(body_lines)
        else:
            subject = "Offer Letter"
            body = letter.content

        doc.add_heading(subject, level=1)
        doc.add_paragraph("")  # Empty line
        for para in body.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())

    doc.save(file_path)

    # Update letter with relative file path
    letter.file_path = str(file_path.relative_to(Path("backend")))
    session.add(letter)
    session.commit()
    session.refresh(letter)

    return letter